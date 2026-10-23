#!/usr/bin/env python3
"""Generate email notification reference workbook (one sheet per module)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = "/Applications/XAMPP/xamppfiles/htdocs/cpci-business-ops-api/email_notifications_by_module.xlsx"

HEADERS = [
    "Internal $subject",
    "Email Subject Sent",
    "Triggered When",
    "Action By (Role ID)",
    "Action By (Role Name)",
    "Primary Recipient Logic",
    "Fallback Recipient Logic",
    "Intended Recipient Role",
    "Notes",
]

MODULES = {
    "Requests": [
        [
            "request_for_approval",
            "Request for Approval: Request #{id}",
            "New request created by Sales AE",
            "19, 4",
            "Sales Account Executive, Admin",
            "immediate_head of creator",
            "creator email",
            "20 - Sales Manager",
            "",
        ],
        [
            "request_for_costing",
            "Request for Costing: Request #{id}",
            "SM/SD creates request (auto-approved on insert)",
            "20, 21",
            "Sales Manager, Sales Director",
            "SM: BDO (32) + immediate_head To/CC; SD: BDO (32)",
            "immediate_head then creator email",
            "32 - BDO (+ SD for SM create)",
            "Auto-approval sets approved_by on insert",
        ],
        [
            "request_for_costing",
            "Request for Costing: Request #{id}",
            "Request status changed to approved",
            "19, 20, 4",
            "Sales AE, Sales Manager, Admin",
            "All users with role 32 (BDO)",
            "approver immediate_head then approver email",
            "32 - Business Development Officer",
            "",
        ],
    ],
    "Costings": [
        [
            "costing_for_approval_by_sales_manager",
            "Costing for Approval by Sales Manager: Costing #{id}",
            "BDO/Business Ops creates costing",
            "22, 32, 4",
            "Business Ops, BDO, Admin",
            "immediate_head of request.added_by",
            "creator immediate_head then creator email",
            "20 - Sales Manager",
            "",
        ],
        [
            "costing_for_approval_by_sales_director",
            "Costing for Approval by Sales Director: Costing #{id}",
            "SM creates costing (auto SM-approved, pending SD) OR SM approves costing",
            "20, 4",
            "Sales Manager, Admin",
            "All users with role 21 (Sales Director)",
            "requester immediate_head then requester email",
            "21 - Sales Director",
            "SM create auto-stamps approved_by_sales_manager",
        ],
        [
            "quotation_for_approval_by_sales",
            "Quotation for Approval by Sales: Quotation #{id}",
            "SD approves BDO-created costing; quotation auto-created",
            "21, 4",
            "Sales Director, Admin",
            "Email of request.added_by",
            "requester immediate_head then requester email",
            "19 - Sales Account Executive",
            "Normal costing approval path",
        ],
        [
            "contracts_for_approval_by_finance_manager",
            "Contract for Approval by Finance Manager: Quotation #{id}",
            "SD creates costing OR SD approves SM-created costing",
            "21, 20",
            "Sales Director, Sales Manager",
            "All users with role 23 (Finance Manager)",
            "requester immediate_head then requester email",
            "23 - Finance Manager",
            "Quotation auto-stamped pending FM approval",
        ],
    ],
    "Quotations": [
        [
            "contracts_for_approval_by_sales_manager",
            "Contract for Approval by Sales Manager: Quotation #{id}",
            "Sales AE approves contract",
            "19, 4",
            "Sales Account Executive, Admin",
            "immediate_head of quotation.added_by",
            "approver immediate_head then approver email",
            "20 - Sales Manager",
            "",
        ],
        [
            "contracts_for_approval_by_sales_director",
            "Contract for Approval by Sales Director: Quotation #{id}",
            "SM approves contract",
            "20, 4",
            "Sales Manager, Admin",
            "All users with role 21",
            "approver immediate_head then approver email",
            "21 - Sales Director",
            "",
        ],
        [
            "contracts_for_approval_by_finance_manager",
            "Contract for Approval by Finance Manager: Quotation #{id}",
            "SD approves contract",
            "21, 4",
            "Sales Director, Admin",
            "All users with role 23",
            "approver immediate_head then approver email",
            "23 - Finance Manager",
            "",
        ],
        [
            "contracts_for_moa",
            "Contract for MOA: Quotation #{id}",
            "FM approves contract (final)",
            "23, 4",
            "Finance Manager, Admin",
            "Email of request.added_by",
            "approver immediate_head then approver email",
            "Request creator / sales chain",
            "",
        ],
    ],
    "Endorsements": [
        [
            "endorsement_for_approval",
            "Endorsement for Approval: Endorsement #{id}",
            "New endorsement created",
            "19, 20",
            "Sales AE, Sales Manager",
            "immediate_head of creator",
            "creator email",
            "20 - Sales Manager (typical)",
            "",
        ],
        [
            "endorsement_revision_for_approval",
            "Endorsement Revision for Approval: Endorsement #{id}",
            "Endorsement revision submitted",
            "19, 20",
            "Sales AE, Sales Manager",
            "immediate_head of submitter",
            "submitter email",
            "20 - Sales Manager (typical)",
            "",
        ],
        [
            "endorsement_approved",
            "Endorsement Approved: Endorsement #{id}",
            "Endorsement status changed to approved",
            "20, 21, 4",
            "Sales Manager, Sales Director, Admin",
            "endorsement.corporate@clinicaprime.com.ph",
            "immediate_head of requester",
            "Corporate endorsement inbox",
            "",
        ],
        [
            "endorsement_revision_approved",
            "Endorsement Revision Approved: Endorsement #{id}",
            "Revision approved or status change while already approved",
            "20, 21, 4",
            "Sales Manager, Sales Director, Admin",
            "endorsement.corporate@clinicaprime.com.ph",
            "immediate_head of requester",
            "Corporate endorsement inbox",
            "",
        ],
        [
            "upload_patients",
            "Upload Patients: Endorsement #{id}",
            "First patient upload",
            "Any",
            "Any authenticated user",
            "All users with role 40 (IT) To/CC",
            "endorsement.corporate@clinicaprime.com.ph",
            "40 - IT",
            "Only on first upload",
        ],
    ],
}


def style_header(ws, row=1, col_count=len(HEADERS)):
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def autosize_columns(ws, min_width=12, max_width=48):
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        length = max(len(str(cell.value or "")) for cell in column_cells)
        width = min(max(length + 2, min_width), max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_module_sheet(wb, sheet_name, rows):
    ws = wb.create_sheet(title=sheet_name)
    ws.append(HEADERS)
    style_header(ws)

    wrap = Alignment(vertical="top", wrap_text=True)
    thin = Side(style="thin", color="EEEEEE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    alt_fill = PatternFill("solid", fgColor="F2F7FB")

    for r_idx, row in enumerate(rows, start=2):
        ws.append(row)
        fill = alt_fill if r_idx % 2 == 0 else None
        for c_idx in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.alignment = wrap
            cell.border = border
            if fill:
                cell.fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    autosize_columns(ws)


def main():
    wb = Workbook()
    wb.remove(wb.active)

    for module_name, rows in MODULES.items():
        write_module_sheet(wb, module_name, rows)

    wb.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
