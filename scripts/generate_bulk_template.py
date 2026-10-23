#!/usr/bin/env python3
"""Generate static request_bulk_template.xlsx for encoder offline use.

Mirrors sheets from Requests::download_bulk_template.
Consultations sheet gets headers only (API download prefills EMR specializations).
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "templates" / "request_bulk_template.xlsx"

REQUEST_HEADERS = [
    "customer_account_id",
    "reference_no",
    "request_type",
    "priority",
    "due_date",
    "contact_person",
    "contact_number",
    "email",
    "delivery_address",
    "payment_terms",
    "grand_total",
    "date",
    "sae_in_charge_id",
    "company_id",
    "mode_of_payment",
    "preferred_branch_id",
    "hmo_account_name",
    "estimated_headcount",
    "number_of_days",
    "shift",
    "with_masterlist",
    "with_walk_in",
    "transaction_type",
    "male_35_plus",
    "female_35_plus",
    "classification",
    "service_duration_date_from",
    "service_duration_date_to",
    "rdt_requester",
    "signatory",
    "signatory_designation",
    "remarks",
    "guaranteed_headcount",
    "mobile_site",
    "email_result",
    "evaluation_data_source",
    "estimated_patient_per_shift",
    "type_cost",
    "mode_of_releasing",
]

COSTING_HEADERS = [
    "package_name",
    "service_package",
    "other_service_package",
    "ape_date",
    "hours_per_day",
    "pre_negotiated_rate",
    "admin_fee",
    "grand_total",
    "remarks",
    "service_duration_date_from",
    "service_duration_date_to",
]

QUOTATION_HEADERS = [
    "version_no",
    "effectivity_date_from",
    "effectivity_date_to",
    "quotation_no",
    "remarks",
    "note",
    "signatory",
    "signatory_designation",
    "grand_total",
]

ENDORSEMENT_HEADERS = [
    "branch_id",
    "mobile_branch_id",
    "type",
    "endorsement_date",
    "transaction_type",
    "address",
    "preferred_branch_id",
    "mode_of_releasing",
    "results_tat",
    "email",
    "time",
    "head_count",
    "payment_terms",
    "date_from",
    "date_to",
    "contract_date_from",
    "contract_date_to",
    "package_name",
    "effectivity_date_from",
    "effectivity_date_to",
    "service_duration_date_from",
    "service_duration_date_to",
    "remarks",
    "note",
]

SERVICES_HEADERS = [
    "type",
    "type_id",
    "type_price",
    "type_cost",
    "branch_id",
    "charge_to",
    "charge_to_ids",
    "head_count",
    "discount_amount",
    "margin_amount",
    "is_guaranteed",
    "pre_negotiated_rate",
    "margin",
    "cost",
    "pro_rate_price",
    "discount",
    "sales",
    "admin_fee",
    "direct_cost",
    "manpower_cost",
    "mobilization_cost",
    "total_cost",
    "estimated_margin",
    "total_amount",
    "theoretical_cost",
]

DETAILS_HEADERS = [
    "lab_test_id",
    "branch_id",
    "quantity",
    "unit_price",
    "total_amount",
    "discount",
    "pro_rate_price",
    "margin",
    "cost",
]

SHIFTS_HEADERS = [
    "name",
    "start_shift",
    "end_shift",
    "registration_cut_off",
    "specimen_submission_cut_off",
    "total_hours",
    "break",
    "ingress",
]

CONSULTATIONS_HEADERS = [
    "specialization_id",
    "specialization_name",
    "company_id",
    "fee_amount",
    "doctor_share",
]


def write_header_sheet(wb, title, headers, single_data_row):
    ws = wb.create_sheet(title)
    ws.append(headers)
    if single_data_row:
        ws.append([""] * len(headers))
    ws.freeze_panes = "A2"
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(headers[i - 1]) + 2)
    return ws


def main():
    wb = Workbook()

    instructions = wb.active
    instructions.title = "Instructions"
    for line in [
        ["Request Bulk Upload Template"],
        [""],
        ["1. Fill the Request sheet (1 data row). Required-ish: company_id, customer_account_id, date, sae_in_charge_id."],
        ["2. Fill Costing / Quotation / Endorsement sheets for entity-specific fields (1 data row each). Shared fields blank on Costing are copied from Request."],
        ["3. Add rows on Services, Details, Shifts as needed. Leave unused rows blank."],
        ["4. Consultations: specialization_id / specialization_name come from doctors/get_specilization."],
        ["   This static file has Consultation headers only. Prefer GET /requests/download_bulk_template for EMR-prefilled IDs."],
        ["   Fill company_id, fee_amount, doctor_share only for rows you need. Rows with empty fee_amount are skipped on upload."],
        ["   Do not change specialization_id / specialization_name when using the API-prefilled template."],
        ["5. Upload via POST /requests/upload_bulk with multipart field \"file\"."],
        ["6. Upload creates one approved Request → Costing → Quotation → Endorsement chain."],
    ]:
        instructions.append(line)
    instructions.column_dimensions["A"].width = 120

    write_header_sheet(wb, "Request", REQUEST_HEADERS, True)
    write_header_sheet(wb, "Costing", COSTING_HEADERS, True)
    write_header_sheet(wb, "Quotation", QUOTATION_HEADERS, True)
    write_header_sheet(wb, "Endorsement", ENDORSEMENT_HEADERS, True)
    write_header_sheet(wb, "Services", SERVICES_HEADERS, False)
    write_header_sheet(wb, "Details", DETAILS_HEADERS, False)
    write_header_sheet(wb, "Shifts", SHIFTS_HEADERS, False)

    consult = write_header_sheet(wb, "Consultations", CONSULTATIONS_HEADERS, False)
    # Placeholder empty row for encoder guidance
    consult.append(["", "", "", "", ""])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
